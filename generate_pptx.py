from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import base64, io

prs = Presentation()
# A4 サイズ
prs.slide_width = Emu(7560000)   # 210mm
prs.slide_height = Emu(10692000) # 297mm

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

W = 7560000
H = 10692000
MM = 36000  # 1mm in EMU

# === ヘッダー ===
header = slide.shapes.add_shape(
    1, 0, 0, W, int(38 * MM)
)
header.fill.solid()
header.fill.fore_color.rgb = RGBColor(0x0d, 0x94, 0x88)
header.line.fill.background()

tf = header.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(16)
run = p.add_run()
run.text = "肥満症治療薬の保険適応をチェック"
run.font.size = Pt(36)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

from pptx.oxml.ns import qn
# サブタイトル
p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(4)
run2 = p2.add_run()
run2.text = "ウゴービ® / ゼップバウンド® があなたに使えるか確認してみましょう"
run2.font.size = Pt(16)
run2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# === イントロ ===
intro_top = int(42 * MM)
intro = slide.shapes.add_shape(
    1, int(8 * MM), intro_top, W - int(16 * MM), int(28 * MM)
)
intro.fill.solid()
intro.fill.fore_color.rgb = RGBColor(0xF0, 0xFD, 0xF4)
intro.line.color.rgb = RGBColor(0xBB, 0xF7, 0xD0)
intro.line.width = Pt(2)

tf = intro.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(8)
run = p.add_run()
run.text = "肥満症は「治療できる病気」です。"
run.font.size = Pt(24)
run.font.bold = True
run.font.color.rgb = RGBColor(0x16, 0x65, 0x34)

p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(4)
run2 = p2.add_run()
run2.text = "一定の条件を満たす方には、\n肥満症治療薬が保険適用で使用できます。"
run2.font.size = Pt(18)
run2.font.color.rgb = RGBColor(0x16, 0x65, 0x34)

# === セクションタイトル ===
sec_top = int(74 * MM)
sec = slide.shapes.add_textbox(int(8 * MM), sec_top, W - int(16 * MM), int(12 * MM))
tf = sec.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "保険適応の主な条件"
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = RGBColor(0x0d, 0x94, 0x88)
# 下線代わりの線
line = slide.shapes.add_shape(
    1, int(8 * MM), sec_top + int(11 * MM), W - int(16 * MM), int(1 * MM)
)
line.fill.solid()
line.fill.fore_color.rgb = RGBColor(0x0d, 0x94, 0x88)
line.line.fill.background()

# === 4つの条件カード ===
conditions = [
    ("1", "BMI 27以上", "身長と体重から計算される\nBMIが27以上であること"),
    ("2", "対象疾患の治療中", "高血圧・脂質異常症・\n2型糖尿病のいずれかで\n薬物治療中"),
    ("3", "BMI別の追加条件", "BMI 35以上：条件1・2を\n満たせば対象\nBMI 27〜35：健康障害が\n2つ以上必要"),
    ("4", "生活習慣改善の実施", "食事・運動療法を行っても\n十分な効果が得られて\nいないこと"),
]

card_w = int(92 * MM)
card_h = int(48 * MM)
gap = int(8 * MM)
margin_x = int(8 * MM)
card_top = int(90 * MM)

for i, (num, title, desc) in enumerate(conditions):
    col = i % 2
    row = i // 2
    x = margin_x + col * (card_w + gap)
    y = card_top + row * (card_h + gap)

    card = slide.shapes.add_shape(1, x, y, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
    card.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)
    card.line.width = Pt(1)

    # 番号丸
    circle = slide.shapes.add_shape(
        9, x + int(3 * MM), y + int(3 * MM), int(9 * MM), int(9 * MM)  # oval
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = RGBColor(0x0d, 0x94, 0x88)
    circle.line.fill.background()
    tf = circle.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = num
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # タイトル
    ttl = slide.shapes.add_textbox(x + int(14 * MM), y + int(3 * MM), card_w - int(17 * MM), int(10 * MM))
    tf = ttl.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

    # 説明
    dtb = slide.shapes.add_textbox(x + int(4 * MM), y + int(15 * MM), card_w - int(8 * MM), card_h - int(18 * MM))
    tf = dtb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = desc
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

# === QRセクション ===
qr_top = card_top + 2 * (card_h + gap) + int(4 * MM)
qr_box = slide.shapes.add_shape(
    1, int(8 * MM), qr_top, W - int(16 * MM), int(50 * MM)
)
qr_box.fill.solid()
qr_box.fill.fore_color.rgb = RGBColor(0xF0, 0xF9, 0xFF)
qr_box.line.color.rgb = RGBColor(0x0d, 0x94, 0x88)
qr_box.line.width = Pt(3)

# QR画像
qr_b64 = "iVBORw0KGgoAAAANSUhEUgAAAlgAAAJYCAYAAAC+ZpjcAAAAAklEQVR4AewaftIAABF1SURBVO3BgZEcAW4EsCbr8k+5rQz8muVrvD4A0z8CAMCZDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABw6icfmpnwPdrmG81Mnmqbt8xMnmqbt8xMnmqbT8xMnmqbT8xMnmqbp2Ymn2ibN8xMfqO2eWpmwvdom6c2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcOonL2ob/s7M5C0zE/5O2zw1M+Hfaps3tM0nZiZPtc1TbcO/1Tb8nZnJGzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKmffLGZyTdqm2/UNm+ZmTw1M/lGbfOJmclTbfPUzOQtM5NPtM0bZiafaJvfZmbylrb5RjOTb9Q232gDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAAp34C/8DM5BNt89TMhH+nbb7VzOQbzUyeapunZiZvaZtPzEzgv2kDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAAp34C/0DbfGJm8lTb/EZt89TM5Km2+cTM5Km2+UZt85aZyVNt84mZyVMzk0+0Dfw3bQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUz/5Ym3DvzMzeaptfqOZyVNt84mZyVNt89TM5C0zk7e0zVtmJk+1zVva5i0zk6fa5hu1Df/OBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE795EUzE/i/bGbyibZ5ambyibZ5ambyVNt8YmbyVNt8YmbyhpnJJ9rmqZnJU23ziZnJU23zG81M+A4bAABObQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODU9I/Af2Bm8pa24e/MTL5R2zw1M/lGbfONZiafaBv4/2oDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAAp37CX5uZ/EZtw9+ZmTzVNm9pm6dmJm9pm0/MTH6bmck3mpm8pW2empm8pW2+0czkE23zhg0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc+smLZiZvaZun2uZbzUyempl8o7b5jdrmqZnJU23zlpnJJ9rmqZnJb9M2n5iZPNU2b5mZPNU2b5mZ8O9sAAA4tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBT0z/ygZnJU23zG81MvlHbPDUzeUvbPDUz+UTb/DYzk7e0zTeamfBvtc1TM5On2uZbzUze0DbfaAMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnfvLFZiZvaJtPtM03mpm8pW3e0DafmJl8o7Z5qm0+MTN5y8zkqbZ5S9s8NTN5qm0+MTN5y8wE/ps2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGr6Rz4wM3lL2zw1M3mqbT4xM+HfaZu3zEze0jZvmJl8q7Z5w8zkG7XNt5qZ/DZt85aZyTdqm6c2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCp6R95ycyEv9M2b5mZPNU2n5iZvKFt3jIzeUvbPDUz+VZt84aZySfa5g0zE/5O27xlZvKJtnlqZvKWtnnDBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE795EMzE/5O2zw1M/lE23yjtnlqZvKWmclb2ua3aZtPzEyeapun2uYTM5M3tM0nZia/zczkE23zVNt8Ymbyhrb5RhsAAE5tAAA4tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NRPPtQ2T81MPtE2T81MnmqbT8xMnmqbT8xMnmqbp2Ymb2mbp2Ymn2ibbzQzeapt3jIz4Xdom99oZvJU23yibZ6amfw2GwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATm0AADg1/SMfmJk81TafmJk81TZPzUw+0TZPzUze0jZvmZn8Nm3ziZnJU23z1MzkW7XNG2Ym36ht3jIz+UTbvGFm8om2ecvM5Ldpm6c2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGr6R77UzOSptnnLzOSptnnLzOQbtc1bZiZPtc0nZiZvaJtPzEyeaptPzEze0DZvmZnwd9rmG81MPtE2/Oc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpn3xoZvJU23yibd4wM3nLzOQtbfPUzOQTbfPUzOQtbfON2oa/0za/Tdt8YmbylrZ5ambyVNu8pW34dzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABw6icvmpl8om2empl8o7b5xMzkDW3zlrZ5y8zkqbb5RNs8NTP5RjOTT7TNN5qZ/DZt84mZyVNt85aZyVNt85aZyTdqm6c2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpn3yxmclTbfPUzOQTbfPUzOQTbfPUzOSptvnEzOQbtc1bZiZPtc03mpl8Ymby27TNW9rmLW3z1MzkLW3Dd9gAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqZ98sbZ5ambyjdrmLW3zjdrmqZnJJ2YmT7XNN5qZfKJtvlHbPDUz+UTbPDUzeUvbPDUz+UTbvKFt3jIz+UTbPNU2v80GAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATv3kRW3zlrb5RjOTT7TNbzMzeaptPjEzeWpm8om2eWpm8hu1zVMzk6fa5hMzkze0zbeamfw2bfOJmclTbfPbbAAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUz/5pWYmb2mbp9rmEzOTp9rmqZnJW9rmLW3z27TNJ2YmT7XNJ2Ym36htvtHM5Km2+cTM5BvNTJ5qm7fMTJ5qm2+0AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFPTP/KBmclv0zZvmZnwd9rmqZnJJ9rmqZnJJ9rmqZnJU23DvzUzeaptnpqZfKu2ecPM5BNt841mJt+obZ7aAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKmf8FXa5i0zk280M3nLzOSptvnEzOSptnlqZsLfaZtPtM1TM5On2uZbzUyeaptvNDP5Rm3zjTYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKmffKht3jIzeaptnpqZfKJtnpqZvKVt+HdmJr9R27xlZvJU2zw1M/lGM5O3tM1bZiZPtc0nZiZPtc0nZiZPtc1vswEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBT0z/yC81MvlHbvGVm8lTbfGJm8o3a5qmZCX+nbT4xM/lt2uYtM5On2uYbzUw+0TbfaGbyVNt8YmbyVNs8tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATv3kQzOTb9Q232hm8om2+UZt89TM5Km2+cTM5Ldpm0/MTJ6amXyibX6bmclTbfOJtnlqZvIbzUze0jbfqG3esAEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBT0z8C/4GZyVva5g0zk0+0zVtmJk+1zVtmJr9N23xiZvJU2zw1M3lL23xiZvJU23yjmclb2ua32QAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpn3xoZsL3aJs3tM0nZiZPtc1bZiZvaRv+nbb5RjOTt7TNbzMz+UTbfKOZyVNt8402AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpn7yobfg7MxP+zszkN5qZPNU2/J2ZyVva5hvNTH6btvlWbfPUzOSpmclb2uapDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJz6yRebmXyjtuF3aJunZia/Udt8o5nJN5qZPNU2n2ibt8xMnpqZ/EYzkze0zSdmJm/YAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnfgL/obZ5S9vwd9qGvzMzeaptvlHbPDUz+VZt84aZySfa5hvNTN7SNm/YAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKmfwD8wM/lE2/B3ZiZPtc1TM5NPtM1TM5O3zEze0jbfqG3eMjN5qm3eMjN5qm2+0czkE23zhg0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc+skXaxv438xMfqO2eWpm8paZyTdqm6dmJm+ZmTzVNm+ZmXyibeC/aQMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnPrJi2Ym/A5t84mZyVNt89TM5BNt843a5qmZyVva5hMzk280M+HfmZm8pW3eMjP5RjOTp9rmqQ0AAKc2AACc2gAAcGoDAMCpDQAApzYAAJzaAABwagMAwKkNAACnNgAAnNoAAHBqAwDAqQ0AAKc2AACcmv4RAADObAAAOLUBAODUBgCAUxsAAE5tAAA4tQEA4NQGAIBTGwAATm0AADi1AQDg1AYAgFMbAABObQAAOLUBAODU/wCZ8uSpOv1GFwAAAABJRU5ErkJggg=="
qr_data = base64.b64decode(qr_b64)
qr_stream = io.BytesIO(qr_data)
qr_img = slide.shapes.add_picture(qr_stream, int(14 * MM), qr_top + int(5 * MM), int(40 * MM), int(40 * MM))

# QRテキスト
qr_txt = slide.shapes.add_textbox(int(60 * MM), qr_top + int(4 * MM), int(130 * MM), int(42 * MM))
tf = qr_txt.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run = p.add_run()
run.text = "スマホで今すぐチェック"
run.font.size = Pt(24)
run.font.bold = True
run.font.color.rgb = RGBColor(0x0d, 0x94, 0x88)

p2 = tf.add_paragraph()
p2.space_before = Pt(6)
run2 = p2.add_run()
run2.text = "QRコードを読み取るだけで\n保険適応の対象か確認できます。\n所要時間はわずか1分です。"
run2.font.size = Pt(16)
run2.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

p3 = tf.add_paragraph()
p3.space_before = Pt(4)
run3 = p3.add_run()
run3.text = "https://obesity-checker-tau.vercel.app"
run3.font.size = Pt(10)
run3.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

# === 注意書き ===
note_top = qr_top + int(54 * MM)
note = slide.shapes.add_shape(
    1, int(8 * MM), note_top, W - int(16 * MM), int(20 * MM)
)
note.fill.solid()
note.fill.fore_color.rgb = RGBColor(0xFF, 0xFB, 0xEB)
note.line.color.rgb = RGBColor(0xFD, 0xE6, 0x8A)
note.line.width = Pt(1)

tf = note.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run = p.add_run()
run.text = "※ このチェックは簡易的なセルフチェックであり、医療行為ではありません。実際の診断・処方は、肥満症治療の教育認定施設にて専門医が行います。チェック結果をもとに、かかりつけ医にご相談ください。"
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)

# === フッター ===
footer_top = note_top + int(22 * MM)
ft = slide.shapes.add_textbox(int(8 * MM), footer_top, W - int(16 * MM), int(8 * MM))
tf = ft.text_frame
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run = p.add_run()
run.text = "日本肥満学会・日本肥満症治療学会「肥満症治療薬の安全・適正使用に関するステートメント」（2025年4月改訂）に基づく"
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(0xB0, 0xB0, 0xB0)

out = "/Users/iwamototatsuya/Desktop/Claude-Projects/obesity-checker/flyer.pptx"
prs.save(out)
print(f"Saved: {out}")
